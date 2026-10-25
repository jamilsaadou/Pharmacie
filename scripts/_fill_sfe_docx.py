import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

D="output/homologation-ab-propre-20261022"
res={r['key']:r for r in json.load(open(D+"/resultats.json")) if r['status']=='certifiee'}

def info(key):
    r=res[key]; return r['request']['rn'], r['totals']['total'], r['security']['codeSECeFDGI']

# test -> list of (caption, image_path)
def cap(mode,key,extra=""):
    num,total,code=info(key)
    return f"Mode {mode} — Facture n° {num} — Total {total:,} FCFA".replace(","," ")+f" — Code SECeF/DGI : {code}"+extra

plan={
 1:[(cap("T.T.C.","01-ttc"),f"{D}/png/01-ttc.png"),(cap("H.T.","01-ht"),f"{D}/png/01-ht.png")],
 3:[(cap("T.T.C.","03-ttc"),f"{D}/png/03-ttc.png"),(cap("H.T.","03-ht"),f"{D}/png/03-ht.png")],
 9:[(cap("T.T.C.","09-ttc"),f"{D}/png/09-ttc.png"),(cap("H.T.","09-ht"),f"{D}/png/09-ht.png")],
 10:[(cap("T.T.C.","10-ttc"),f"{D}/png/10-ttc.png"),(cap("H.T.","10-ht"),f"{D}/png/10-ht.png")],
 25:[(cap("T.T.C.","25-ttc"),f"{D}/png/25-ttc.png"),(cap("H.T.","25-ht"),f"{D}/png/25-ht.png")],
 26:[(cap("T.T.C.","26-ttc")+" — page 1/2",f"{D}/png/26-ttc-1.png"),
     ("Mode T.T.C. — page 2/2 : éléments de sécurité et code QR",f"{D}/png/26-ttc-2.png"),
     (cap("H.T.","26-ht")+" — page 1/2",f"{D}/png/26-ht-1.png"),
     ("Mode H.T. — page 2/2 : éléments de sécurité et code QR",f"{D}/png/26-ht-2.png")],
}

d=docx.Document("tmp/sfe-docx/original.docx")
markers=[p for p in d.paragraphs if 'INSÉRER LA PHOTO' in p.text]
assert len(markers)==26, f"attendu 26 marqueurs, trouvé {len(markers)}"

def new_caption(text):
    p=d.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pf=p.paragraph_format
    pf.page_break_before=True   # chaque photo sur sa propre page
    pf.keep_with_next=True      # la légende reste collée à son image
    pf.space_after=Pt(4)
    r=p.add_run(text); r.bold=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x19,0x19,0x19)
    return p._p

def new_image(path):
    d.add_picture(path, width=Inches(5.6))
    p=d.paragraphs[-1]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together=True
    return p._p

for test,items in plan.items():
    marker=markers[test-1]
    els=[]
    for caption,img in items:
        els.append(new_caption(caption))
        els.append(new_image(img))
    for el in reversed(els):
        marker._p.addnext(el)
    print(f"test #{test}: {len(items)} photo(s) insérée(s)")

out="tmp/sfe-docx/SFE cas de test - rempli.docx"
d.save(out)
print("Enregistré :",out)
